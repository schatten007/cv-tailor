"""Export real documents, inspect text, and stage versioned application files."""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from html import escape
from pathlib import Path

import reportlab
from docx import Document
from docx.shared import Mm, Pt
from docx.table import Table
from PIL import Image
from pypdf import PdfReader, PdfWriter
from pypdf.errors import PdfReadError
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from careerkit.contracts import SCHEMAS, Json, local_path, read_json, safe_id, validate, write_json

PLACEHOLDER = re.compile(
    r"\[(?:confirm|metric needed|name|company|date|email|phone|insert\b[^\]]*|todo\b[^\]]*)\]",
    re.IGNORECASE,
)
EXTENSIONS = {".pdf", ".docx", ".png", ".jpg", ".jpeg"}


def _read_pdf(path: Path) -> PdfReader:
    try:
        return PdfReader(path)
    except PdfReadError as error:
        raise ValueError(f"Unreadable PDF: {path.name}") from error


def private_output(path: Path) -> Path:
    """Keep generated candidate files outside the installed public skill checkout."""
    resolved = path.resolve()
    if resolved.is_relative_to(SCHEMAS.parent):
        raise ValueError("Choose a career workspace outside the skill repository")
    return resolved


def digest(path: Path) -> str:
    """Calculate a file's full SHA-256 without loading it all into memory."""
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def final_text(text: str) -> None:
    """Reject known draft placeholders and empty final documents."""
    if not text.strip():
        raise ValueError("No extractable text; supply text or OCR before using this document")
    match = PLACEHOLDER.search(text)
    if match:
        raise ValueError(f"Resolve draft placeholder before exporting or staging: {match.group()}")


def inspect_text(path: Path) -> str:
    """Extract supported source text in document order; do not pretend to perform OCR."""
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".json"}:
        return path.read_text(encoding="utf-8-sig")
    if suffix == ".pdf":
        reader = _read_pdf(path)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDF: provide an unlocked copy")
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".docx":
        document = Document(str(path))
        parts: list[str] = []
        seen: set[str] = set()
        for section in document.sections:
            for header in (section.header, section.first_page_header, section.even_page_header):
                key = str(header.part.partname)
                if key not in seen:
                    parts.extend(p.text for p in header.paragraphs)
                    seen.add(key)
        for block in document.iter_inner_content():
            if isinstance(block, Table):
                parts.extend(" | ".join(cell.text for cell in row.cells) for row in block.rows)
            else:
                parts.append(block.text)
        text = "\n".join(parts)
    else:
        raise ValueError(
            "Unsupported extraction format; provide text or an available OCR/extraction tool"
        )
    if not text.strip():
        raise ValueError("No extractable text; supply text or OCR before using this document")
    return text


def document_lines(data: Json) -> list[tuple[str, str]]:
    """Flatten a validated document into a single deterministic reading order."""
    validate("document", data)
    lines = [("title", data["title"])]
    lines.extend(("body", line) for line in data.get("contact", []))
    for section in data["sections"]:
        if section.get("heading"):
            lines.append(("heading", section["heading"]))
        lines.extend(("body", paragraph) for paragraph in section.get("paragraphs", []))
        for entry in section.get("entries", []):
            lines.append(("entry", entry["heading"]))
            for key in ("subheading", "dates"):
                if entry.get(key):
                    lines.append(("body", entry[key]))
            lines.extend(("bullet", bullet) for bullet in entry.get("bullets", []))
    final_text("\n".join(text for _, text in lines))
    return lines


def _docx_bytes(lines: list[tuple[str, str]]) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(18)
    section.left_margin = section.right_margin = Mm(20)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for kind, text in lines:
        if kind in {"title", "heading"}:
            document.add_heading(text, level=0 if kind == "title" else 1)
        else:
            paragraph = document.add_paragraph(style="List Bullet" if kind == "bullet" else None)
            run = paragraph.add_run(text)
            run.bold = kind == "entry"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_bytes(lines: list[tuple[str, str]], font: Path | None) -> bytes:
    font_path = font or Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"
    pdf_font = TTFont("CVTailor", str(font_path))
    characters = {
        ord(character) for _, text in lines for character in text if not character.isspace()
    }
    missing = sorted(characters - set(pdf_font.face.charToGlyph))
    if missing:
        raise ValueError(
            "PDF font lacks glyphs: "
            + ", ".join(f"U+{code:04X}" for code in missing)
            + "; supply --font with a suitable TTF"
        )
    pdfmetrics.registerFont(pdf_font)
    styles = {
        "title": ParagraphStyle(
            "title", fontName="CVTailor", fontSize=18, leading=23, spaceAfter=8
        ),
        "heading": ParagraphStyle(
            "heading",
            fontName="CVTailor",
            fontSize=13,
            leading=17,
            spaceBefore=10,
            spaceAfter=5,
            keepWithNext=True,
        ),
        "entry": ParagraphStyle(
            "entry", fontName="CVTailor", fontSize=11, leading=15, spaceBefore=5, keepWithNext=True
        ),
        "body": ParagraphStyle("body", fontName="CVTailor", fontSize=11, leading=15, spaceAfter=4),
        "bullet": ParagraphStyle(
            "bullet",
            fontName="CVTailor",
            fontSize=11,
            leading=15,
            leftIndent=10,
            firstLineIndent=-7,
            spaceAfter=3,
        ),
    }
    story = []
    for kind, text in lines:
        plain = ("- " if kind == "bullet" else "") + text
        story.append(Paragraph(escape(plain).replace("\n", "<br/>"), styles[kind]))
    story.append(Spacer(1, 1))
    output = io.BytesIO()
    SimpleDocTemplate(
        output, pagesize=A4, leftMargin=56, rightMargin=56, topMargin=51, bottomMargin=51
    ).build(story)
    return output.getvalue()


def export_document(
    data: Json, output_base: Path, format_name: str, font: Path | None = None
) -> list[Path]:
    """Export requested real files without overwriting earlier document versions."""
    lines = document_lines(data)
    base = private_output(output_base)
    if format_name not in {"docx", "pdf", "both"}:
        raise ValueError("Format must be docx, pdf, or both")
    formats = ["docx", "pdf"] if format_name == "both" else [format_name]
    outputs = [Path(str(base) + "." + extension) for extension in formats]
    if any(path.exists() for path in outputs):
        raise ValueError("An output already exists; choose a new version name")
    # Build all formats first so font/schema failures create no partial exports.
    payloads = [
        _docx_bytes(lines) if extension == "docx" else _pdf_bytes(lines, font)
        for extension in formats
    ]
    base.parent.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    try:
        for path, payload in zip(outputs, payloads, strict=True):
            with path.open("xb") as stream:
                stream.write(payload)
            written.append(path)
    except OSError:
        for path in written:
            path.unlink(missing_ok=True)
        raise
    return outputs


def merge_pdfs(inputs: list[Path], output: Path) -> int:
    """Merge explicitly selected PDFs in order without replacing source files."""
    destination = private_output(output)
    if not inputs or output.suffix.lower() != ".pdf":
        raise ValueError("Provide PDF inputs and a .pdf output")
    if destination.exists():
        raise ValueError("Output exists; choose a new combined-document filename")
    writer = PdfWriter()
    for path in inputs:
        reader = _read_pdf(path)
        if reader.is_encrypted or not reader.pages:
            raise ValueError("Cannot combine encrypted or empty PDFs")
        writer.append(reader)
    count = len(writer.pages)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("xb") as stream:
        writer.write(stream)
    writer.close()
    return count


def _check_file(path: Path, kind: str) -> None:
    suffix = path.suffix.lower()
    if suffix not in EXTENSIONS or not path.is_file() or not path.stat().st_size:
        raise ValueError("Select a nonempty PDF, DOCX, PNG, or JPEG file")
    if suffix == ".pdf":
        reader = _read_pdf(path)
        if reader.is_encrypted or not reader.pages:
            raise ValueError("Select an unlocked, nonempty PDF")
    elif suffix == ".docx":
        with zipfile.ZipFile(path) as archive:
            if "word/document.xml" not in archive.namelist():
                raise ValueError("File is not a DOCX document")
        Document(str(path))
    else:
        with Image.open(path) as image:
            image.verify()
        if kind != "supporting":
            raise ValueError("CVs and cover letters need text-based PDF or DOCX, not images")
    if kind in {"cv", "cover-letter"}:
        final_text(inspect_text(path))


def load_manifest(workspace: Path) -> Json:
    """Read a workspace manifest or an empty manifest before initial staging."""
    path = workspace / "documents.json"
    data = read_json(path) if path.exists() else {"schema_version": 1, "documents": {}}
    validate("manifest", data)
    return data


def stage_document(
    workspace: Path,
    source: Path,
    identifier: str,
    kind: str,
    language: str,
    *,
    replace: bool = False,
) -> Json:
    """Copy one selected file into a versioned workspace path and record its hash."""
    root = private_output(workspace)
    safe_id(identifier)
    if kind not in {"cv", "cover-letter", "supporting"} or len(language) < 2:
        raise ValueError("Supply a valid document kind and language")
    _check_file(source, kind)
    manifest = load_manifest(root)
    checksum = digest(source)
    previous = manifest["documents"].get(identifier)
    if previous and previous["sha256"] != checksum and not replace:
        raise ValueError("Document ID already staged; use --replace to select a new version")
    relative = f"documents/{identifier}-{checksum[:16]}{source.suffix.lower()}"
    destination = local_path(root, relative)
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists():
        if digest(destination) != checksum:
            raise ValueError("Existing staged version is modified; choose another workspace")
    else:
        with destination.open("xb") as stream:
            stream.write(source.read_bytes())
    record = {
        "kind": kind,
        "path": relative,
        "sha256": checksum,
        "size_bytes": destination.stat().st_size,
        "language": language,
    }
    manifest["documents"][identifier] = record
    validate("manifest", manifest)
    write_json(root / "documents.json", manifest, replace=True)
    return record


def verify_documents(workspace: Path, identifiers: list[str]) -> Json:
    """Resolve selected upload paths, checking containment, size, hash, and content."""
    manifest = load_manifest(workspace)
    records: Json = {}
    if len(set(identifiers)) != len(identifiers):
        raise ValueError("Duplicate selected document IDs")
    for identifier in identifiers:
        safe_id(identifier)
        record = manifest["documents"].get(identifier)
        if record is None:
            raise ValueError(f"No staged document: {identifier}")
        path = local_path(workspace, record["path"])
        if (
            not path.is_file()
            or path.stat().st_size != record["size_bytes"]
            or digest(path) != record["sha256"]
        ):
            raise ValueError(f"Missing or changed staged document: {identifier}")
        _check_file(path, record["kind"])
        records[identifier] = {**record, "absolute_path": str(path)}
    return records
