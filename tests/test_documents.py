"""Test real export/readback, Unicode, staging, and upload identity checks."""

import copy
import zipfile
from pathlib import Path

import pytest
from docx import Document
from PIL import Image
from pypdf import PdfReader, PdfWriter

from careerkit.contracts import SCHEMAS, Json, local_path, read_json, write_json
from careerkit.documents import (
    export_document,
    final_text,
    inspect_text,
    load_manifest,
    merge_pdfs,
    private_output,
    stage_document,
    verify_documents,
)


def test_exports_real_single_column_files_with_unicode(tmp_path: Path, document: Json) -> None:
    document["title"] = "Jörg Müller"
    document["sections"].append(
        {
            "heading": "Additional information",
            "paragraphs": ["<literal text> & Grüße, €20 per hour"],
        }
    )
    paths = export_document(document, tmp_path / "unicode-cv", "both")
    for path in paths:
        text = inspect_text(path)
        assert "Jörg Müller" in text and "Grüße" in text and "€20" in text
        assert "<literal text>" in text
        assert text.index("Education") < text.index("Projects") < text.index("Skills")
    docx = Document(str(paths[0]))
    assert not docx.tables
    assert not any(p.text for p in docx.sections[0].header.paragraphs)
    pdf = PdfReader(paths[1])
    assert len(pdf.pages) == 1
    assert float(pdf.pages[0].mediabox.width) == pytest.approx(595.27, abs=0.1)
    assert float(pdf.pages[0].mediabox.height) == pytest.approx(841.89, abs=0.1)


def test_bad_glyph_fails_without_partial_outputs(tmp_path: Path, document: Json) -> None:
    document["title"] = "Example 🛰"
    with pytest.raises(ValueError, match="glyphs"):
        export_document(document, tmp_path / "unsupported", "both")
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize(
    "placeholder", ["[confirm]", "[metric needed]", "[Name]", "[insert salary]", "[TODO: dates]"]
)
def test_final_exports_reject_placeholders(
    tmp_path: Path, document: Json, placeholder: str
) -> None:
    document["title"] = placeholder
    with pytest.raises(ValueError, match="placeholder"):
        export_document(document, tmp_path / "invalid", "docx")


def test_export_keeps_versions_and_rejects_invalid_format(tmp_path: Path, document: Json) -> None:
    export_document(document, tmp_path / "cv", "docx")
    with pytest.raises(ValueError, match="already exists"):
        export_document(document, tmp_path / "cv", "both")
    assert not (tmp_path / "cv.pdf").exists()
    with pytest.raises(ValueError, match="Format"):
        export_document(document, tmp_path / "cv", "png")
    with pytest.raises(ValueError, match="outside"):
        private_output(SCHEMAS.parent / "personal-cv.pdf")


def test_source_inspection_handles_tables_headers_and_bom(tmp_path: Path) -> None:
    path = tmp_path / "source.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Candidate name"
    document.add_paragraph("Experience")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "2024 - 2026"
    table.cell(0, 1).text = "Engineer"
    document.save(str(path))
    extracted = inspect_text(path)
    assert (
        extracted.index("Candidate name")
        < extracted.index("Experience")
        < extracted.index("Engineer")
    )
    text = tmp_path / "notes.md"
    text.write_text("\ufeffCandidate supplied notes", encoding="utf-8")
    assert inspect_text(text) == "Candidate supplied notes"
    with pytest.raises(ValueError, match="Unsupported"):
        inspect_text(tmp_path / "image.png")


def test_scan_encrypted_and_corrupt_pdf_are_not_treated_as_text(tmp_path: Path) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    scan = tmp_path / "scan.pdf"
    writer.write(scan)
    with pytest.raises(ValueError, match="OCR"):
        inspect_text(scan)
    with pytest.raises(ValueError, match="extractable"):
        final_text("  ")
    writer.encrypt("synthetic-test-password")
    locked = tmp_path / "locked.pdf"
    writer.write(locked)
    with pytest.raises(ValueError, match="Encrypted"):
        inspect_text(locked)
    with pytest.raises(ValueError, match="unlocked"):
        stage_document(tmp_path / "app", locked, "certificate", "supporting", "en")
    bad = tmp_path / "renamed.pdf"
    bad.write_text("Markdown is not a PDF", encoding="utf-8")
    with pytest.raises(ValueError, match="Unreadable PDF"):
        inspect_text(bad)


def test_merge_selected_pdfs_in_order(tmp_path: Path, document: Json) -> None:
    first = export_document(document, tmp_path / "cv", "pdf")[0]
    letter = copy.deepcopy(document)
    letter["title"] = "Application letter"
    second = export_document(letter, tmp_path / "letter", "pdf")[0]
    out = tmp_path / "combined.pdf"
    assert merge_pdfs([first, second], out) == 2
    text = inspect_text(out)
    assert text.index("Alex Example") < text.index("Application letter")
    with pytest.raises(ValueError, match="Output exists"):
        merge_pdfs([first], out)
    with pytest.raises(ValueError, match="Provide PDF"):
        merge_pdfs([], tmp_path / "empty.pdf")
    with pytest.raises(ValueError, match="Provide PDF"):
        merge_pdfs([first], tmp_path / "not-pdf.docx")
    empty = tmp_path / "zero.pdf"
    PdfWriter().write(empty)
    with pytest.raises(ValueError, match="empty"):
        merge_pdfs([empty], tmp_path / "failed.pdf")


def test_staging_is_idempotent_and_versions_are_preserved(workspace: Path, document: Json) -> None:
    first = load_manifest(workspace)["documents"]["cv"]
    old = local_path(workspace, first["path"])
    assert stage_document(workspace, old, "cv", "cv", "en") == first
    document["title"] = "Alex Example - revised"
    new = export_document(document, workspace.parent / "revised", "pdf")[0]
    with pytest.raises(ValueError, match="--replace"):
        stage_document(workspace, new, "cv", "cv", "en")
    changed = stage_document(workspace, new, "cv", "cv", "en", replace=True)
    assert old.exists() and changed["path"] != first["path"]
    assert verify_documents(workspace, ["cv"])["cv"]["sha256"] == changed["sha256"]


@pytest.mark.parametrize("mode", ["modified", "missing", "traversal", "unknown", "duplicate"])
def test_manifest_cannot_upload_wrong_files(workspace: Path, mode: str) -> None:
    manifest = load_manifest(workspace)
    path = local_path(workspace, manifest["documents"]["cv"]["path"])
    ids = ["cv"]
    if mode == "modified":
        path.write_bytes(b"changed file")
    elif mode == "missing":
        path.unlink()
    elif mode == "traversal":
        manifest["documents"]["cv"]["path"] = "../unrelated.pdf"
        write_json(workspace / "documents.json", manifest, replace=True)
    elif mode == "unknown":
        ids = ["letter"]
    else:
        ids = ["cv", "cv"]
    with pytest.raises(ValueError):
        verify_documents(workspace, ids)


def test_supporting_images_are_allowed_but_not_image_cvs(tmp_path: Path) -> None:
    image = tmp_path / "certificate.png"
    Image.new("RGB", (20, 20), "white").save(image)
    stage_document(tmp_path / "app", image, "certificate", "supporting", "de")
    assert verify_documents(tmp_path / "app", ["certificate"])
    with pytest.raises(ValueError, match="not images"):
        stage_document(tmp_path / "app", image, "cv", "cv", "en")
    with pytest.raises(ValueError, match="valid document kind"):
        stage_document(tmp_path / "app", image, "other", "unknown", "en")
    bad = tmp_path / "empty.pdf"
    bad.touch()
    with pytest.raises(ValueError, match="nonempty"):
        stage_document(tmp_path / "app", bad, "empty", "supporting", "en")


def test_docx_file_type_is_checked_and_modified_version_not_overwritten(
    tmp_path: Path, workspace: Path
) -> None:
    fake = tmp_path / "fake.docx"
    with zipfile.ZipFile(fake, "w") as archive:
        archive.writestr("not-a-document.txt", "hello")
    with pytest.raises(ValueError, match="not a DOCX"):
        stage_document(workspace, fake, "letter", "cover-letter", "en")
    original = workspace.parent / "cv.pdf"
    manifest = read_json(workspace / "documents.json")
    staged = local_path(workspace, manifest["documents"]["cv"]["path"])
    staged.write_bytes(b"modified")
    with pytest.raises(ValueError, match="staged version is modified"):
        stage_document(workspace, original, "cv", "cv", "en")
